import docx
import re

doc = docx.Document()

# Page margins: Top 2.54cm, Bottom 2.54cm, Left 3.17cm, Right 2.54cm
for s in doc.sections:
    s.top_margin = docx.shared.Cm(2.54)
    s.bottom_margin = docx.shared.Cm(2.54)
    s.left_margin = docx.shared.Cm(3.17)
    s.right_margin = docx.shared.Cm(2.54)

# Base styling
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = docx.shared.Pt(12)
normal.paragraph_format.line_spacing = 1.5

def add_formatted_runs(paragraph, text):
    """Parses markdown bold (**text**), italics (*text*), and code (`code`) into native Word runs."""
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith('*') and token.endswith('*'):
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        elif token.startswith('`') and token.endswith('`'):
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Courier New'
            r.font.size = docx.shared.Pt(10)
        else:
            paragraph.add_run(token)

src_md = '/home/cantroll/MATUNDUSMS/docs/Deborah_Mbula_Test_Plan.md'
with open(src_md, 'r') as f:
    lines = f.readlines()

in_table = False
table_headers = []
table_rows = []

def flush_table():
    global in_table, table_headers, table_rows
    if in_table and table_headers:
        t = doc.add_table(rows=1, cols=len(table_headers))
        t.style = 'Light Grid Accent 1'
        t.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        
        # Header row
        for i, h in enumerate(table_headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            clean_h = h.strip()
            if clean_h.startswith('**') and clean_h.endswith('**'):
                clean_h = clean_h[2:-2]
            r = p.add_run(clean_h)
            r.bold = True
            
        # Data rows
        for row in table_rows:
            rc = t.add_row().cells
            for i, val in enumerate(row):
                if i < len(rc):
                    rc[i].text = ""
                    p = rc[i].paragraphs[0]
                    add_formatted_runs(p, val.strip())
        doc.add_paragraph()
    in_table = False
    table_headers = []
    table_rows = []

for line in lines:
    raw = line.rstrip()
    stripped = raw.strip()

    # Table parsing
    if stripped.startswith('|') and stripped.endswith('|'):
        cols = [c.strip() for c in stripped.split('|')[1:-1]]
        if all(set(c).issubset({'-', ' ', ':'}) for c in cols if c):
            continue  # Separator line
        if not in_table:
            in_table = True
            table_headers = cols
        else:
            table_rows.append(cols)
        continue
    else:
        if in_table:
            flush_table()

    if not stripped:
        continue

    if stripped.startswith('# '):
        p = doc.add_paragraph()
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        clean_text = stripped[2:].replace('**', '')
        r = p.add_run(clean_text.upper())
        r.bold = True
        r.font.size = docx.shared.Pt(16)
    elif stripped.startswith('## '):
        clean_h = stripped[3:].strip().replace('**', '')
        doc.add_heading(clean_h, level=1)
    elif stripped.startswith('### '):
        clean_h = stripped[4:].strip().replace('**', '')
        doc.add_heading(clean_h, level=2)
    elif stripped.startswith('#### '):
        clean_h = stripped[5:].strip().replace('**', '')
        doc.add_heading(clean_h, level=3)
    elif stripped.startswith('- ') or stripped.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        add_formatted_runs(p, stripped[2:])
    elif stripped.startswith('---'):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = docx.shared.Pt(6)
    elif stripped.startswith('```'):
        continue
    else:
        p = doc.add_paragraph()
        add_formatted_runs(p, stripped)

if in_table:
    flush_table()

docx_out = '/home/cantroll/MATUNDUSMS/docs/Deborah_Mbula_Test_Plan.docx'
doc.save(docx_out)
print('DOCX generated cleanly with native bold formatting:', docx_out)
